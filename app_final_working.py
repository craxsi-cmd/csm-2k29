from pathlib import Path
import io
import json
from datetime import date, datetime

import cv2
import numpy as np
import pandas as pd
import streamlit as st
from insightface.app import FaceAnalysis
from PIL import Image
from fpdf import FPDF
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ============================================================
# CONFIGURATION
# ============================================================

APP_TITLE = "ClassAttend AI"
APP_SUBTITLE = "Smart classroom attendance using computer vision"

BASE_DIR = Path(__file__).parent
DATA_DIR = BASE_DIR / "data"
PHOTO_DIR = DATA_DIR / "student_photos"

STUDENTS_FILE = DATA_DIR / "students.csv"
EMBEDDINGS_FILE = DATA_DIR / "face_embeddings.json"
EXCEL_FILE = DATA_DIR / "attendance_master.xlsx"

MATCH_THRESHOLD = 0.35
DETECTION_THRESHOLD = 0.35
DETECTION_SIZE = (1280, 1280)

DATA_DIR.mkdir(parents=True, exist_ok=True)
PHOTO_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# PAGE CONFIGURATION AND STYLE
# ============================================================

st.set_page_config(
    page_title=APP_TITLE,
    page_icon="🎓",
    layout="wide",
    initial_sidebar_state="expanded",
)

st.markdown(
    """
    <style>
        .main {
            background-color: #f5f7fb;
        }

        .block-container {
            padding-top: 2rem;
            padding-bottom: 3rem;
            max-width: 1400px;
        }

        [data-testid="stSidebar"] {
            background-color: #102a43;
        }

        [data-testid="stSidebar"] * {
            color: white;
        }

        h1, h2, h3 {
            color: #102a43;
        }

        .hero {
            background: linear-gradient(
                135deg,
                #102a43,
                #1f6f8b
            );
            padding: 2rem;
            border-radius: 18px;
            color: white;
            margin-bottom: 1.5rem;
        }

        .hero h1 {
            color: white;
            margin-bottom: 0.3rem;
        }

        .hero p {
            color: #d9f0ff;
            font-size: 1.05rem;
        }

        .metric-card {
            background: white;
            padding: 1rem;
            border-radius: 14px;
            border: 1px solid #d9e2ec;
            box-shadow: 0 2px 8px rgba(16, 42, 67, 0.08);
        }

        .info-card {
            background: white;
            padding: 1.3rem;
            border-radius: 14px;
            border-left: 5px solid #1f8a70;
            box-shadow: 0 2px 8px rgba(16, 42, 67, 0.08);
            margin: 0.8rem 0;
        }

        .stButton > button {
            width: 100%;
            border-radius: 9px;
            min-height: 2.7rem;
            font-weight: 600;
        }

        .stDownloadButton > button {
            width: 100%;
            border-radius: 9px;
            min-height: 2.7rem;
            font-weight: 600;
        }

        footer {
            visibility: hidden;
        }
    </style>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# MODEL
# ============================================================

@st.cache_resource
def load_face_model():
    model = FaceAnalysis(
        name="buffalo_l",
        providers=["CPUExecutionProvider"],
    )

    model.prepare(
        ctx_id=0,
        det_thresh=DETECTION_THRESHOLD,
        det_size=DETECTION_SIZE,
    )

    return model


# ============================================================
# DATA FUNCTIONS
# ============================================================

def load_students():
    if not STUDENTS_FILE.exists():
        return pd.DataFrame(
            columns=[
                "Roll Number",
                "Name",
                "Photo File",
            ]
        )

    return pd.read_csv(
        STUDENTS_FILE,
        dtype=str,
    ).fillna("")


def load_embeddings():
    if not EMBEDDINGS_FILE.exists():
        return {}

    try:
        with open(
            EMBEDDINGS_FILE,
            "r",
            encoding="utf-8",
        ) as file:
            return json.load(file)
    except Exception:
        return {}


def save_embeddings(embeddings):
    with open(
        EMBEDDINGS_FILE,
        "w",
        encoding="utf-8",
    ) as file:
        json.dump(
            embeddings,
            file,
            indent=2,
        )


def image_from_upload(uploaded_file):
    image_bytes = uploaded_file.getvalue()

    image_array = np.frombuffer(
        image_bytes,
        dtype=np.uint8,
    )

    return cv2.imdecode(
        image_array,
        cv2.IMREAD_COLOR,
    )


def get_image_dimensions(image):
    if image is None:
        return 0, 0

    height, width = image.shape[:2]
    return width, height


# ============================================================
# FACE FUNCTIONS
# ============================================================

def extract_embedding(uploaded_photo):
    image = image_from_upload(uploaded_photo)

    if image is None:
        return None, "The image could not be read."

    model = load_face_model()
    faces = model.get(image)

    if len(faces) == 0:
        return None, "No face was detected."

    if len(faces) > 1:
        return None, (
            "Multiple faces detected. Upload one clear photo "
            "containing one student."
        )

    embedding = faces[0].normed_embedding

    return embedding.astype(float).tolist(), None


def get_best_match(query_embedding, embeddings):
    best_roll_number = None
    best_similarity = -1.0

    query = np.array(
        query_embedding,
        dtype=np.float32,
    )

    query_norm = np.linalg.norm(query)

    if query_norm == 0:
        return None, 0.0

    query = query / query_norm

    for roll_number, saved_embedding in embeddings.items():
        saved = np.array(
            saved_embedding,
            dtype=np.float32,
        )

        saved_norm = np.linalg.norm(saved)

        if saved_norm == 0:
            continue

        saved = saved / saved_norm

        similarity = float(
            np.dot(query, saved)
        )

        if similarity > best_similarity:
            best_similarity = similarity
            best_roll_number = roll_number

    return best_roll_number, best_similarity


# ============================================================
# STUDENT REGISTRATION
# ============================================================

def save_student(roll_number, name, uploaded_photo):
    students = load_students()

    if roll_number in students["Roll Number"].tolist():
        return False, "This roll number already exists."

    safe_name = "".join(
        character if character.isalnum() else "_"
        for character in name
    )

    photo_path = PHOTO_DIR / f"{roll_number}_{safe_name}.jpg"

    image = Image.open(
        io.BytesIO(uploaded_photo.getvalue())
    ).convert("RGB")

    image.save(
        photo_path,
        "JPEG",
        quality=95,
    )

    new_student = pd.DataFrame(
        [
            {
                "Roll Number": roll_number,
                "Name": name,
                "Photo File": str(
                    photo_path.relative_to(BASE_DIR)
                ),
            }
        ]
    )

    students = pd.concat(
        [
            students,
            new_student,
        ],
        ignore_index=True,
    )

    students.to_csv(
        STUDENTS_FILE,
        index=False,
    )

    return True, "Student registered successfully."


# ============================================================
# EXCEL REGISTER
# ============================================================

def style_excel_header(worksheet):
    for cell in worksheet[1]:
        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )


def update_excel_register():
    students = load_students()

    if EXCEL_FILE.exists():
        workbook = load_workbook(EXCEL_FILE)
        worksheet = workbook["Attendance Register"]

        existing_roll_numbers = {
            str(
                worksheet.cell(
                    row=row,
                    column=1,
                ).value
            )
            for row in range(
                2,
                worksheet.max_row + 1,
            )
        }

        for _, student in students.iterrows():
            roll_number = str(student["Roll Number"])

            if roll_number not in existing_roll_numbers:
                worksheet.append(
                    [
                        student["Roll Number"],
                        student["Name"],
                    ]
                )

    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Attendance Register"

        worksheet.append(
            [
                "Roll Number",
                "Student Name",
            ]
        )

        for _, student in students.iterrows():
            worksheet.append(
                [
                    student["Roll Number"],
                    student["Name"],
                ]
            )

    style_excel_header(worksheet)

    worksheet.column_dimensions["A"].width = 16
    worksheet.column_dimensions["B"].width = 30
    worksheet.freeze_panes = "C2"
    worksheet.auto_filter.ref = worksheet.dimensions

    workbook.save(EXCEL_FILE)


def add_attendance_column(present_roll_numbers):
    today = date.today().isoformat()

    update_excel_register()

    workbook = load_workbook(EXCEL_FILE)
    worksheet = workbook["Attendance Register"]

    headers = [
        worksheet.cell(
            row=1,
            column=column,
        ).value
        for column in range(
            1,
            worksheet.max_column + 1,
        )
    ]

    if today in headers:
        date_column = headers.index(today) + 1
    else:
        date_column = worksheet.max_column + 1

        worksheet.cell(
            row=1,
            column=date_column,
        ).value = today

    for row in range(2, worksheet.max_row + 1):
        roll_number = str(
            worksheet.cell(
                row=row,
                column=1,
            ).value
        )

        if roll_number in present_roll_numbers:
            status = "✓ Present"
        else:
            status = "✗ Absent"

        cell = worksheet.cell(
            row=row,
            column=date_column,
        )

        cell.value = status
        cell.alignment = Alignment(
            horizontal="center"
        )

        if "Present" in status:
            cell.font = Font(
                color="008000",
                bold=True,
            )
        else:
            cell.font = Font(
                color="C00000",
                bold=True,
            )

    header_cell = worksheet.cell(
        row=1,
        column=date_column,
    )

    header_cell.font = Font(
        bold=True,
        color="FFFFFF",
    )

    header_cell.fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_cell.alignment = Alignment(
        horizontal="center",
    )

    worksheet.column_dimensions[
        header_cell.column_letter
    ].width = 16

    workbook.save(EXCEL_FILE)

    return today


# ============================================================
# PDF EXPORT
# ============================================================

def generate_attendance_pdf():
    update_excel_register()

    workbook = load_workbook(EXCEL_FILE)
    worksheet = workbook["Attendance Register"]

    rows = list(
        worksheet.iter_rows(
            values_only=True
        )
    )

    if not rows:
        return b""

    headers = [
        str(value) if value is not None else ""
        for value in rows[0]
    ]

    pdf = FPDF(
        orientation="L",
        unit="mm",
        format="A4",
    )

    pdf.set_auto_page_break(
        auto=True,
        margin=12,
    )

    pdf.add_page()

    pdf.set_font(
        "Helvetica",
        "B",
        16,
    )

    pdf.cell(
        0,
        10,
        "CLASSROOM ATTENDANCE REGISTER",
        new_x="LMARGIN",
        new_y="NEXT",
        align="C",
    )

    pdf.set_font(
        "Helvetica",
        "",
        9,
    )

    pdf.cell(
        0,
        7,
        f"Generated on: {date.today().isoformat()}",
        new_x="LMARGIN",
        new_y="NEXT",
        align="C",
    )

    pdf.ln(4)

    page_width = 297
    margin = 10
    usable_width = page_width - (2 * margin)

    total_columns = len(headers)
    date_columns = max(
        0,
        total_columns - 2,
    )

    roll_width = 22
    name_width = 65

    if date_columns:
        date_width = min(
            24,
            max(
                14,
                (
                    usable_width
                    - roll_width
                    - name_width
                )
                / date_columns,
            ),
        )
    else:
        date_width = 24

    pdf.set_font(
        "Helvetica",
        "B",
        7,
    )

    pdf.set_fill_color(
        31,
        78,
        121,
    )

    pdf.set_text_color(
        255,
        255,
        255,
    )

    pdf.cell(
        roll_width,
        10,
        "Roll No.",
        border=1,
        fill=True,
        align="C",
    )

    pdf.cell(
        name_width,
        10,
        "Student Name",
        border=1,
        fill=True,
        align="C",
    )

    for header in headers[2:]:
        pdf.cell(
            date_width,
            10,
            header,
            border=1,
            fill=True,
            align="C",
        )

    pdf.ln()

    pdf.set_text_color(
        0,
        0,
        0,
    )

    pdf.set_font(
        "Helvetica",
        "",
        7,
    )

    for row in rows[1:]:
        roll_number = str(
            row[0]
            if len(row) > 0 and row[0] is not None
            else ""
        )

        student_name = str(
            row[1]
            if len(row) > 1 and row[1] is not None
            else ""
        )

        pdf.cell(
            roll_width,
            8,
            roll_number,
            border=1,
            align="C",
        )

        pdf.cell(
            name_width,
            8,
            student_name[:32],
            border=1,
        )

        for value in row[2:]:
            status = str(
                value if value is not None else ""
            )

            if "Present" in status:
                short_status = "P"
                pdf.set_text_color(
                    0,
                    120,
                    0,
                )

            elif "Absent" in status:
                short_status = "A"
                pdf.set_text_color(
                    180,
                    0,
                    0,
                )

            else:
                short_status = "-"
                pdf.set_text_color(
                    0,
                    0,
                    0,
                )

            pdf.cell(
                date_width,
                8,
                short_status,
                border=1,
                align="C",
            )

        pdf.set_text_color(
            0,
            0,
            0,
        )

        pdf.ln()

    pdf.ln(5)

    pdf.set_font(
        "Helvetica",
        "",
        8,
    )

    pdf.cell(
        0,
        6,
        "P = Present    A = Absent",
        new_x="LMARGIN",
        new_y="NEXT",
    )

    return bytes(
        pdf.output()
    )


# ============================================================
# UI HELPERS
# ============================================================

def show_hero(title, subtitle):
    st.markdown(
        f"""
        <div class="hero">
            <h1>{title}</h1>
            <p>{subtitle}</p>
        </div>
        """,
        unsafe_allow_html=True,
    )


def show_metric(label, value):
    st.markdown(
        f"""
        <div class="metric-card">
            <small>{label}</small>
            <h2>{value}</h2>
        </div>
        """,
        unsafe_allow_html=True,
    )


# ============================================================
# APPLICATION
# ============================================================

show_hero(
    APP_TITLE,
    APP_SUBTITLE,
)

st.sidebar.markdown(
    "## 🎓 ClassAttend AI"
)

st.sidebar.caption(
    "Teacher attendance workspace"
)

page = st.sidebar.radio(
    "Navigation",
    [
        "Dashboard",
        "Register Student",
        "Mark Attendance",
        "Face Recognition Test",
        "Attendance Register",
    ],
)

students = load_students()
embeddings = load_embeddings()


# ============================================================
# DASHBOARD
# ============================================================

if page == "Dashboard":
    st.header("Dashboard")

    first, second, third, fourth = st.columns(4)

    with first:
        show_metric(
            "Registered Students",
            len(students),
        )

    with second:
        show_metric(
            "Face Profiles",
            len(embeddings),
        )

    with third:
        show_metric(
            "Register File",
            "Ready" if EXCEL_FILE.exists() else "New",
        )

    with fourth:
        show_metric(
            "Recognition Mode",
            "CPU AI",
        )

    st.markdown(
        """
        <div class="info-card">
            <h3>Teacher workflow</h3>
            <p>
                Register students once, upload one clear classroom photo,
                analyze the faces, and download the updated attendance
                register as Excel or PDF.
            </p>
        </div>
        """,
        unsafe_allow_html=True,
    )

    st.subheader("Getting started")

    step_one, step_two, step_three = st.columns(3)

    with step_one:
        st.info(
            "1. Register each student with a clear photo."
        )

    with step_two:
        st.info(
            "2. Upload the classroom photo in Mark Attendance."
        )

    with step_three:
        st.info(
            "3. Download the Excel or PDF register."
        )

    if not students.empty:
        st.subheader("Registered Students")

        st.dataframe(
            students[
                [
                    "Roll Number",
                    "Name",
                ]
            ],
            use_container_width=True,
            hide_index=True,
        )


# ============================================================
# REGISTER STUDENT
# ============================================================

elif page == "Register Student":
    st.header("Register Student")

    st.info(
        "Use one clear photo containing exactly one face."
    )

    roll_number = st.text_input(
        "Roll number",
        placeholder="Example: 101",
    )

    name = st.text_input(
        "Full name",
        placeholder="Example: Rahul Sharma",
    )

    uploaded_photo = st.file_uploader(
        "Upload student photo",
        type=[
            "jpg",
            "jpeg",
            "png",
        ],
        help=(
            "Use a sharp, front-facing image with good lighting."
        ),
    )

    if uploaded_photo is not None:
        st.image(
            uploaded_photo,
            caption="Selected student photo",
            width=300,
        )

    if st.button(
        "Register Student",
        type="primary",
    ):
        roll_number = roll_number.strip()
        name = name.strip()

        if (
            not roll_number
            or not name
            or uploaded_photo is None
        ):
            st.error(
                "Enter roll number, name, and upload a photo."
            )

        else:
            current_students = load_students()

            if roll_number in current_students[
                "Roll Number"
            ].tolist():
                st.error(
                    "This roll number already exists."
                )

            else:
                with st.spinner(
                    "Detecting face and creating embedding..."
                ):
                    embedding, error = extract_embedding(
                        uploaded_photo
                    )

                if error:
                    st.error(error)

                else:
                    success, message = save_student(
                        roll_number,
                        name,
                        uploaded_photo,
                    )

                    if success:
                        current_embeddings = load_embeddings()

                        current_embeddings[
                            roll_number
                        ] = embedding

                        save_embeddings(
                            current_embeddings
                        )

                        update_excel_register()

                        st.success(
                            "Student and face profile registered."
                        )

                        st.rerun()

                    else:
                        st.error(message)

    st.subheader("Registered Students")

    if students.empty:
        st.info(
            "No students registered yet."
        )

    else:
        st.dataframe(
            students[
                [
                    "Roll Number",
                    "Name",
                ]
            ],
            use_container_width=True,
            hide_index=True,
        )


# ============================================================
# FACE RECOGNITION TEST
# ============================================================

elif page == "Face Recognition Test":
    st.header("Face Recognition Test")

    if students.empty:
        st.warning(
            "Register at least one student first."
        )

    else:
        st.write(
            "Test recognition before marking a classroom attendance session."
        )

        test_photo = st.file_uploader(
            "Upload test photo",
            type=[
                "jpg",
                "jpeg",
                "png",
            ],
            key="test_photo",
        )

        if test_photo is not None:
            st.image(
                test_photo,
                caption="Test photo",
                width=500,
            )

            if st.button(
                "Recognize Face",
                type="primary",
            ):
                image = image_from_upload(
                    test_photo
                )

                if image is None:
                    st.error(
                        "The uploaded image could not be read."
                    )

                else:
                    model = load_face_model()
                    faces = model.get(image)

                    st.write(
                        f"Detected faces: {len(faces)}"
                    )

                    if len(faces) == 0:
                        st.warning(
                            "No face detected."
                        )

                    else:
                        results = []

                        for face in faces:
                            face_width = int(
                                face.bbox[2]
                                - face.bbox[0]
                            )

                            face_height = int(
                                face.bbox[3]
                                - face.bbox[1]
                            )

                            roll_number, similarity = (
                                get_best_match(
                                    face.normed_embedding,
                                    embeddings,
                                )
                            )

                            matched = (
                                roll_number is not None
                                and similarity >= MATCH_THRESHOLD
                            )

                            if matched:
                                student = students[
                                    students[
                                        "Roll Number"
                                    ].astype(str)
                                    == str(roll_number)
                                ]

                                if not student.empty:
                                    results.append(
                                        {
                                            "Roll Number": roll_number,
                                            "Name": student.iloc[0][
                                                "Name"
                                            ],
                                            "Similarity": round(
                                                similarity,
                                                3,
                                            ),
                                            "Face Size": (
                                                f"{face_width} x "
                                                f"{face_height}"
                                            ),
                                            "Status": "Recognized",
                                        }
                                    )

                            else:
                                results.append(
                                    {
                                        "Roll Number": "-",
                                        "Name": "Unknown face",
                                        "Similarity": round(
                                            similarity,
                                            3,
                                        ),
                                        "Face Size": (
                                            f"{face_width} x "
                                            f"{face_height}"
                                        ),
                                        "Status": "Review",
                                    }
                                )

                        st.dataframe(
                            pd.DataFrame(results),
                            use_container_width=True,
                            hide_index=True,
                        )


# ============================================================
# MARK ATTENDANCE
# ============================================================

elif page == "Mark Attendance":
    st.header("Mark Attendance")

    st.info(
        "Upload the original high-resolution classroom photo. "
        "Avoid screenshots or compressed images."
    )

    if students.empty:
        st.warning(
            "Register students before marking attendance."
        )

    else:
        classroom_photo = st.file_uploader(
            "Upload classroom photo",
            type=[
                "jpg",
                "jpeg",
                "png",
            ],
            key="classroom_photo",
        )

        if classroom_photo is not None:
            st.image(
                classroom_photo,
                caption="Classroom photo",
                use_container_width=True,
            )

            if st.button(
                "Analyze and Mark Attendance",
                type="primary",
            ):
                with st.spinner(
                    "Detecting and recognizing classroom faces..."
                ):
                    image = image_from_upload(
                        classroom_photo
                    )

                    if image is None:
                        st.error(
                            "The classroom image could not be read."
                        )

                    else:
                        model = load_face_model()
                        faces = model.get(image)

                        current_embeddings = load_embeddings()
                        recognized = []
                        unknown_count = 0
                        diagnostics = []

                        for face in faces:
                            face_width = int(
                                face.bbox[2]
                                - face.bbox[0]
                            )

                            face_height = int(
                                face.bbox[3]
                                - face.bbox[1]
                            )

                            roll_number, similarity = (
                                get_best_match(
                                    face.normed_embedding,
                                    current_embeddings,
                                )
                            )

                            matched = (
                                roll_number is not None
                                and similarity >= MATCH_THRESHOLD
                            )

                            if matched:
                                if roll_number not in recognized:
                                    recognized.append(
                                        roll_number
                                    )

                                status = "Recognized"

                            else:
                                unknown_count += 1
                                status = "Review"

                            diagnostics.append(
                                {
                                    "Roll Number": (
                                        roll_number
                                        if matched
                                        else "-"
                                    ),
                                    "Similarity": round(
                                        similarity,
                                        3,
                                    ),
                                    "Face Size": (
                                        f"{face_width} x "
                                        f"{face_height}"
                                    ),
                                    "Status": status,
                                }
                            )

                        today = add_attendance_column(
                            set(recognized)
                        )

                        st.success(
                            f"Attendance saved for {today}."
                        )

                        first, second, third = st.columns(3)

                        with first:
                            st.metric(
                                "Detected Faces",
                                len(faces),
                            )

                        with second:
                            st.metric(
                                "Recognized Present",
                                len(recognized),
                            )

                        with third:
                            st.metric(
                                "Needs Review",
                                unknown_count,
                            )

                        st.subheader(
                            "Recognition diagnostics"
                        )

                        if diagnostics:
                            st.dataframe(
                                pd.DataFrame(diagnostics),
                                use_container_width=True,
                                hide_index=True,
                            )

                        present_students = students[
                            students[
                                "Roll Number"
                            ].astype(str).isin(recognized)
                        ][
                            [
                                "Roll Number",
                                "Name",
                            ]
                        ]

                        st.subheader(
                            "Present Students"
                        )

                        if present_students.empty:
                            st.warning(
                                "No registered students were recognized."
                            )

                        else:
                            st.dataframe(
                                present_students,
                                use_container_width=True,
                                hide_index=True,
                            )

                        st.info(
                            "All registered students not recognized "
                            "in this photo were marked absent."
                        )


# ============================================================
# ATTENDANCE REGISTER
# ============================================================

elif page == "Attendance Register":
    st.header("Attendance Register")

    update_excel_register()

    if EXCEL_FILE.exists():
        workbook = load_workbook(
            EXCEL_FILE
        )

        worksheet = workbook[
            "Attendance Register"
        ]

        rows = list(
            worksheet.iter_rows(
                values_only=True
            )
        )

        if len(rows) > 1:
            table = pd.DataFrame(
                rows[1:],
                columns=rows[0],
            )

            st.dataframe(
                table,
                use_container_width=True,
                hide_index=True,
            )

        st.subheader("Downloads")

        excel_col, pdf_col = st.columns(2)

        with excel_col:
            with open(
                EXCEL_FILE,
                "rb",
            ) as excel_file:
                st.download_button(
                    label="Download Excel Register",
                    data=excel_file.read(),
                    file_name="attendance_master.xlsx",
                    mime=(
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    ),
                )

        with pdf_col:
            st.download_button(
                label="Download PDF Register",
                data=generate_attendance_pdf(),
                file_name="attendance_register.pdf",
                mime="application/pdf",
            )

    else:
        st.info(
            "Register a student first."
        )


# ============================================================
# FOOTER
# ============================================================

st.markdown("---")

st.caption(
    "ClassAttend AI • Academic prototype • "
    "Face data is stored locally on this computer"
)